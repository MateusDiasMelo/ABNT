"""
Document Processor Service
Processa arquivos DOCX e PDF, extraindo conteúdo e aplicando formatação ABNT.
"""

from docx import Document
import PyPDF2
from pdf2docx import Converter
from pathlib import Path
from typing import Tuple, Optional
import tempfile
import os
from .abnt_formatter import ABNTFormatter


class DocumentProcessor:
    """Processador de documentos DOCX e PDF."""

    @staticmethod
    def process_docx(input_path: str, output_path: str) -> Tuple[Document, int]:
        """
        Processa um arquivo DOCX aplicando formatação ABNT.

        Args:
            input_path: Caminho do arquivo de entrada
            output_path: Caminho do arquivo de saída

        Returns:
            Tupla com (documento formatado, número de páginas)
        """
        # Carrega documento
        doc = Document(input_path)

        # Aplica formatação ABNT
        formatter = ABNTFormatter(doc)
        formatted_doc = formatter.format_document()

        # Conta páginas
        page_count = formatter.count_pages()

        # Salva documento formatado
        formatted_doc.save(output_path)

        return formatted_doc, page_count

    @staticmethod
    def process_pdf(input_path: str, output_path: str) -> Tuple[Document, int]:
        """
        Processa um arquivo PDF convertendo para DOCX e aplicando formatação ABNT.

        Args:
            input_path: Caminho do arquivo PDF de entrada
            output_path: Caminho do arquivo DOCX de saída

        Returns:
            Tupla com (documento formatado, número de páginas)
        """
        # Cria arquivo temporário para conversão
        temp_docx = tempfile.NamedTemporaryFile(suffix='.docx', delete=False)
        temp_docx_path = temp_docx.name
        temp_docx.close()

        try:
            # Converte PDF para DOCX
            cv = Converter(input_path)
            cv.convert(temp_docx_path, start=0, end=None)
            cv.close()

            # Processa o DOCX temporário
            formatted_doc, page_count = DocumentProcessor.process_docx(
                temp_docx_path, output_path
            )

            return formatted_doc, page_count

        finally:
            # Remove arquivo temporário
            if os.path.exists(temp_docx_path):
                os.unlink(temp_docx_path)

    @staticmethod
    def get_pdf_page_count(pdf_path: str) -> int:
        """
        Obtém o número de páginas de um PDF.

        Args:
            pdf_path: Caminho do arquivo PDF

        Returns:
            Número de páginas
        """
        with open(pdf_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            return len(pdf_reader.pages)

    @staticmethod
    def validate_document(file_path: str, extension: str) -> Tuple[bool, Optional[str]]:
        """
        Valida se o documento pode ser processado.

        Args:
            file_path: Caminho do arquivo
            extension: Extensão do arquivo (docx ou pdf)

        Returns:
            Tupla com (é válido, mensagem de erro)
        """
        if not os.path.exists(file_path):
            return False, "Arquivo não encontrado"

        if extension not in ['docx', 'pdf']:
            return False, "Formato de arquivo não suportado"

        # Verifica se o arquivo não está vazio
        if os.path.getsize(file_path) == 0:
            return False, "Arquivo vazio"

        # Tenta abrir o arquivo para validar integridade
        try:
            if extension == 'docx':
                doc = Document(file_path)
                if not doc.paragraphs:
                    return False, "Documento DOCX vazio ou corrompido"
            elif extension == 'pdf':
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    if len(pdf_reader.pages) == 0:
                        return False, "PDF vazio"
        except Exception as e:
            return False, f"Erro ao validar arquivo: {str(e)}"

        return True, None

    @staticmethod
    def extract_text_preview(file_path: str, extension: str, max_chars: int = 500) -> str:
        """
        Extrai uma prévia do texto do documento.

        Args:
            file_path: Caminho do arquivo
            extension: Extensão do arquivo
            max_chars: Número máximo de caracteres

        Returns:
            Prévia do texto
        """
        text = ""

        try:
            if extension == 'docx':
                doc = Document(file_path)
                for para in doc.paragraphs[:10]:  # Primeiros 10 parágrafos
                    text += para.text + "\n"
                    if len(text) >= max_chars:
                        break
            elif extension == 'pdf':
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    for page_num in range(min(2, len(pdf_reader.pages))):  # Primeiras 2 páginas
                        page = pdf_reader.pages[page_num]
                        text += page.extract_text() + "\n"
                        if len(text) >= max_chars:
                            break
        except Exception as e:
            return f"Erro ao extrair texto: {str(e)}"

        return text[:max_chars] + ("..." if len(text) > max_chars else "")


class DocumentAnalyzer:
    """Analisador de documentos para estatísticas e validação."""

    @staticmethod
    def analyze_document(file_path: str, extension: str) -> dict:
        """
        Analisa um documento e retorna estatísticas.

        Args:
            file_path: Caminho do arquivo
            extension: Extensão do arquivo

        Returns:
            Dicionário com estatísticas
        """
        stats = {
            "original_pages": 0,
            "paragraphs": 0,
            "words": 0,
            "characters": 0,
            "has_images": False,
            "has_tables": False,
            "has_references": False,
        }

        try:
            if extension == 'docx':
                doc = Document(file_path)
                stats["paragraphs"] = len(doc.paragraphs)
                stats["has_tables"] = len(doc.tables) > 0

                # Conta palavras e caracteres
                text = " ".join(p.text for p in doc.paragraphs)
                stats["words"] = len(text.split())
                stats["characters"] = len(text)

                # Verifica referências
                for para in doc.paragraphs:
                    if para.text.strip().upper() in ['REFERÊNCIAS', 'REFERENCIAS']:
                        stats["has_references"] = True
                        break

                # Estimativa de páginas originais
                stats["original_pages"] = max(1, stats["characters"] // 1800)

            elif extension == 'pdf':
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    stats["original_pages"] = len(pdf_reader.pages)

                    # Extrai texto para análise
                    text = ""
                    for page in pdf_reader.pages[:50]:  # Limita a 50 páginas
                        text += page.extract_text()

                    stats["words"] = len(text.split())
                    stats["characters"] = len(text)

                    # Verifica referências
                    if 'REFERÊNCIAS' in text.upper() or 'REFERENCIAS' in text.upper():
                        stats["has_references"] = True

        except Exception as e:
            stats["error"] = str(e)

        return stats
